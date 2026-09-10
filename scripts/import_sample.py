"""Read the SCADA sample without modifying the workbook. Always use Time."""
import json
import re
import sys
from pathlib import Path
from datetime import datetime
import openpyxl

source = Path(sys.argv[1])
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = workbook['Sheet1']
headers = [c.value for c in next(sheet.iter_rows())]
required = ['设备编码', '设备名称', '属性编码', '属性名称', '属性采集值', 'Time']
assert all(h in headers for h in required), 'Missing required source columns'
readings = []
for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    if not any(v is not None for v in row):
        continue
    record = dict(zip(headers, row))
    frame = {'A100': 1, 'A101': 2, 'A102': 3, 'A103': 4}[record['属性编码']]
    machine = re.search(r'W1-(\d+)-', record['设备名称']).group(1)
    time = datetime.fromisoformat(str(record['Time'])).strftime('%Y-%m-%d %H:%M:%S')
    readings.append({'id': f'ws1-sheet1-row-{row_num}', 'workshop': 'WS1', 'lineId': f'L{int(machine) // 100}', 'machineId': machine, 'equipmentCode': str(record['设备编码']), 'frame': frame, 'weight': float(record['属性采集值']), 'Time': time, 'productType': None})
readings.sort(key=lambda r: (r['Time'], r['machineId'], r['frame'], r['id']))
payload = {'metadata': {'mode': 'sample', 'source': source.name, 'sheet': 'Sheet1', 'timestampField': 'Time', 'timezone': 'source-local-unspecified', 'from': readings[0]['Time'], 'to': readings[-1]['Time'], 'recordCount': len(readings), 'machineCount': len({r['machineId'] for r in readings}), 'units': 'g', 'productAssignment': 'not-provided'}, 'readings': readings}
destination = Path(__file__).resolve().parent.parent / 'data' / 'sample.json'
destination.parent.mkdir(exist_ok=True)
destination.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(json.dumps(payload['metadata']))
